from pathlib import Path
import re
import zipfile

from docx import Document


ROOT = Path(__file__).resolve().parent
FILES = sorted(ROOT.glob("*.docx"))
EXPECTED = {
    "01_个人陈述_PS_导出模板.docx": ("Personal Statement (PS)", "PERSONAL STATEMENT"),
    "02_学术简历_CV_导出模板.docx": ("Academic Curriculum Vitae (CV)", "CURRICULUM VITAE"),
    "03_推荐信_RL_导出模板.docx": ("Recommendation Letter (RL)", "LETTER OF RECOMMENDATION"),
    "04_研究计划_RP_导出模板.docx": ("Research Proposal (RP)", "RESEARCH PROPOSAL"),
    "05_套磁信_CL_导出模板.docx": ("Cold Email (CL)", "EMAIL CONTENT"),
}
REQUIRED_METADATA = [
    "Document Type:",
    "Student Name:",
    "Student Profile:",
    "Project ID:",
    "Current Institution:",
    "Current Major:",
    "Academic Performance:",
    "Target Institution:",
    "Target Programme:",
    "Application Level:",
    "Application Round:",
    "Intended Entry:",
]

assert len(FILES) == 5, f"Expected 5 DOCX files, found {len(FILES)}"
assert set(path.name for path in FILES) == set(EXPECTED), "Unexpected DOCX filename set"

for path in FILES:
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None, f"Corrupt DOCX package: {path.name}"
        assert "word/document.xml" in package.namelist(), f"Missing document.xml: {path.name}"

    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    full_text = "\n".join(paragraphs)
    document_type, body_heading = EXPECTED[path.name]

    assert len(document.tables) == 0, f"Tables are not allowed: {path.name}"
    assert "STUDENT & APPLICATION INFORMATION" in paragraphs, f"Missing information heading: {path.name}"
    assert body_heading in paragraphs, f"Missing body heading: {path.name}"
    information_index = paragraphs.index("STUDENT & APPLICATION INFORMATION")
    body_heading_index = paragraphs.index(body_heading, information_index + 1)
    assert information_index < body_heading_index, f"Metadata must precede body: {path.name}"
    assert document_type in full_text, f"Missing document type: {path.name}"
    for field in REQUIRED_METADATA:
        assert field in full_text, f"Missing metadata field {field} in {path.name}"

    body_start = body_heading_index + 1
    body_text = "\n".join(paragraphs[body_start:])
    assert re.search(r"[A-Za-z]{4}", body_text), f"No English body content: {path.name}"
    assert not re.search(r"[\u4e00-\u9fff]", body_text), f"Chinese found in English body: {path.name}"

    section = document.sections[0]
    assert abs(section.page_width.cm - 21.0) < 0.05, f"Page width is not A4: {path.name}"
    assert abs(section.page_height.cm - 29.7) < 0.05, f"Page height is not A4: {path.name}"
    assert abs(section.left_margin.cm - 2.5) < 0.05, f"Incorrect left margin: {path.name}"
    assert abs(section.right_margin.cm - 2.5) < 0.05, f"Incorrect right margin: {path.name}"
    assert abs(section.top_margin.cm - 2.2) < 0.05, f"Incorrect top margin: {path.name}"
    assert abs(section.bottom_margin.cm - 2.2) < 0.05, f"Incorrect bottom margin: {path.name}"

    body_paragraphs = [p for p in document.paragraphs[body_start:] if p.text.strip() and not p.text.isupper()]
    line_spacings = [p.paragraph_format.line_spacing for p in body_paragraphs]
    assert line_spacings and all(abs(float(value) - 1.5) < 0.01 for value in line_spacings if value is not None), f"Body line spacing is not 1.5: {path.name}"

    print(f"PASS  {path.name}  {path.stat().st_size:,} bytes  {len(paragraphs)} paragraphs")

cold_email = Document(ROOT / "05_套磁信_CL_导出模板.docx")
cold_paragraphs = [paragraph.text for paragraph in cold_email.paragraphs]
email_start = cold_paragraphs.index("EMAIL CONTENT") + 1
email_text = " ".join(cold_paragraphs[email_start:])
email_word_count = len(re.findall(r"\b[A-Za-z0-9]+(?:[–'-][A-Za-z0-9]+)*\b", email_text))
assert 150 <= email_word_count <= 220, f"Cold email should stay concise; found {email_word_count} words"
print(f"PASS  Cold-email length: {email_word_count} words")
print("All five Word templates passed structural and formatting validation.")
