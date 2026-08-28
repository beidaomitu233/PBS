from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_DIR = Path(__file__).resolve().parent

STUDENT_APPLICATION_INFO = [
    ("Student Name", "Li Xue"),
    ("Student Profile", "Female, 22 years old"),
    ("Contact", "138****8888"),
    ("Project ID", "PRJ20260701001"),
    ("Current Institution", "Peking University"),
    ("Current Major", "Computer Science"),
    ("Academic Performance", "GPA 3.8/4.0; IELTS 7.5; GRE 330"),
    ("Target Institution", "Carnegie Mellon University"),
    ("Target Programme", "Artificial Intelligence"),
    ("Application Level", "Master's Degree"),
    ("Application Round", "Early Decision (ED)"),
    ("Intended Entry", "Fall 2026"),
]

DOCUMENTS = [
    {
        "filename": "01_个人陈述_PS_导出模板.docx",
        "title": "PERSONAL STATEMENT",
        "short": "PS",
        "document_type": "Personal Statement (PS)",
        "body_heading": "PERSONAL STATEMENT",
        "mode": "essay",
        "sections": [
            (
                None,
                [
                    "My interest in artificial intelligence began not with a finished model, but with a question that refused to disappear: why can a system perform impressively on a benchmark yet become unreliable when the data are scarce, noisy, or drawn from a different setting? During my undergraduate study in Computer Science at Peking University, this question developed from curiosity into a clear academic direction. Courses in machine learning, deep learning, algorithms, and natural language processing gave me the theoretical vocabulary to examine it, while research and engineering projects taught me that dependable intelligence requires more than a high headline accuracy.",
                    "The most important step in this development was my work on a national undergraduate innovation project concerning deep-learning-based medical image analysis. As the project lead, I helped define the research problem, coordinated the experimental plan, and implemented core models in PyTorch. Our early results were unstable when the number of labelled samples was reduced, and information from images and clinical text was not consistently aligned. I therefore compared alternative representation-learning strategies, redesigned the evaluation protocol, and documented failure cases rather than reporting only the best run. This experience showed me how careful problem formulation, reproducible experiments, and honest analysis can turn a broad technical ambition into a researchable question.",
                    "My internship at Baidu AI Lab further strengthened this practical perspective. Working on intent classification and slot filling for a dialogue system, I contributed to an improvement in intent-recognition accuracy from 91.2 percent to 94.6 percent before the model was integrated into an online customer-service scenario. The experience required me to balance model quality with latency, maintainability, and the distribution of real user queries. It also taught me to communicate experimental evidence to colleagues with different responsibilities. I came to understand that useful AI research must connect mathematical insight, sound implementation, and the constraints of deployment.",
                    "These experiences have led me to focus on robust multimodal representation learning, especially under limited supervision. I want to investigate how contrastive objectives, attention-based fusion, and prototype-based regularisation can produce representations that remain reliable when one modality is incomplete or the sample distribution shifts. My prior work has prepared me to approach this area with both research discipline and engineering fluency. Two co-authored papers, including an ACL Workshop paper, exposed me to academic writing and peer feedback, while a top-five-percent result in a Kaggle multimodal classification challenge strengthened my ability to iterate independently under time constraints.",
                    "Carnegie Mellon University's Artificial Intelligence programme is a compelling next step because of its combination of rigorous machine-learning foundations, interdisciplinary research, and an environment in which ideas are tested through implementation. I hope to deepen my understanding of representation learning, probabilistic modelling, and trustworthy AI while contributing my experience in multimodal systems and medical data. I am particularly interested in project-based opportunities that connect methodological research with high-stakes applications, where robustness and interpretability are essential rather than optional.",
                    "In the short term, I intend to develop a stronger research framework for studying multimodal learning and to produce work that is reproducible, clearly evaluated, and useful to other researchers. In the longer term, I hope to lead applied research on dependable AI systems for healthcare and other data-sensitive domains. Graduate study at Carnegie Mellon would allow me to refine the questions that emerged from my undergraduate projects, learn from a demanding academic community, and build systems whose performance can be trusted beyond controlled experiments. I would bring to the programme sustained curiosity, careful execution, and a willingness to improve through evidence and collaboration.",
                ],
            )
        ],
    },
    {
        "filename": "02_学术简历_CV_导出模板.docx",
        "title": "ACADEMIC CURRICULUM VITAE",
        "short": "CV",
        "document_type": "Academic Curriculum Vitae (CV)",
        "body_heading": "CURRICULUM VITAE",
        "mode": "sectioned",
        "sections": [
            (
                "PROFILE",
                [
                    "Computer Science undergraduate at Peking University with a GPA of 3.8/4.0 and research interests in multimodal learning, robust representation learning, and medical artificial intelligence. Experienced in research design, deep-learning implementation, experimental evaluation, and technical communication. Proficient in Python and PyTorch, with academic and industry experience developing models for image, text, and dialogue data."
                ],
            ),
            (
                "EDUCATION",
                [
                    "Peking University, Bachelor of Science in Computer Science, September 2022 to June 2026. GPA: 3.8/4.0; academic standing: top 5 percent. Relevant coursework includes Data Structures and Algorithms, Operating Systems, Machine Learning, Deep Learning, Natural Language Processing, Probability and Statistics, and Linear Algebra."
                ],
            ),
            (
                "RESEARCH AND PROJECT EXPERIENCE",
                [
                    "Project Lead, National Undergraduate Innovation Project: Deep-Learning-Based Medical Image Analysis System, September 2025 to present. Defined the research scope, coordinated the experimental plan, and implemented multimodal representation models in PyTorch. Investigated image-text alignment under limited supervision, developed reproducible evaluation pipelines, and analysed model behaviour across different sample sizes and missing-modality settings.",
                    "Participant, Kaggle Multimodal Classification Challenge. Designed and evaluated feature-fusion strategies for image and text inputs, conducted error analysis, and achieved a top-five-percent final ranking. The project strengthened practical skills in data preprocessing, model selection, ablation testing, and experiment tracking."
                ],
            ),
            (
                "INDUSTRY EXPERIENCE",
                [
                    "Natural Language Processing Algorithm Intern, Baidu AI Lab, June 2025 to September 2025. Supported the optimisation of intent classification and slot filling for a production dialogue system. Improved intent-recognition accuracy from 91.2 percent to 94.6 percent through data analysis, model tuning, and targeted evaluation; assisted with integration into an online customer-service workflow."
                ],
            ),
            (
                "PUBLICATIONS AND HONOURS",
                [
                    "Co-author of two research papers, including one paper published through an ACL Workshop. Recipient of the National Scholarship and the University First-Class Scholarship. Holder of one patent related to an undergraduate technology project. Complete publication, patent, and award details should be populated from the verified student record at export time."
                ],
            ),
            (
                "TECHNICAL SKILLS AND LANGUAGES",
                [
                    "Programming and tools: Python, PyTorch, Git, data preprocessing, deep-learning experimentation, and reproducible model evaluation. Languages and tests: Chinese (native), English (IELTS 7.5), and GRE 330."
                ],
            ),
        ],
    },
    {
        "filename": "03_推荐信_RL_导出模板.docx",
        "title": "RECOMMENDATION LETTER",
        "short": "RL",
        "document_type": "Recommendation Letter (RL)",
        "body_heading": "LETTER OF RECOMMENDATION",
        "mode": "letter",
        "sections": [
            (
                None,
                [
                    "To the Admissions Committee,",
                    "I am pleased to recommend Li Xue for admission to the Master's programme in Artificial Intelligence at Carnegie Mellon University. I have known Li through both academic instruction and research supervision in the School of Computer Science at Peking University. Over this period, I have observed her analytical ability, intellectual independence, and professional conduct in settings that required sustained learning as well as practical problem solving.",
                    "Li distinguished herself in advanced coursework by moving beyond the reproduction of standard methods. In machine learning assignments and discussions, she consistently examined why a method worked, what assumptions supported it, and how its behaviour changed when those assumptions were weakened. Her academic record, including a GPA of 3.8/4.0 and a position within the top 5 percent of her cohort, reflects strong preparation. More importantly, her questions and written analyses demonstrate an ability to connect mathematical concepts with evidence from experiments.",
                    "I saw these qualities most clearly while supervising her national undergraduate innovation project on deep-learning-based medical image analysis. Li helped transform a broad topic into a focused investigation of multimodal alignment under limited labelled data. She organised the experimental workflow, implemented core models in PyTorch, and introduced controlled comparisons when initial results proved unstable. Rather than concealing negative outcomes, she traced failure cases and used them to revise the research plan. This response showed maturity, persistence, and the judgement required for graduate-level work.",
                    "Li is also an effective and considerate collaborator. She communicates technical decisions clearly, keeps reliable records, and supports team members when implementation problems arise. During group reviews, she is receptive to criticism and able to distinguish between defending an idea and defending evidence. Her internship experience at Baidu AI Lab, where she contributed to an increase in intent-recognition accuracy from 91.2 percent to 94.6 percent, further demonstrates that she can apply research skills within a demanding engineering environment.",
                    "I believe Li is well prepared for the intellectual rigour and collaborative culture of Carnegie Mellon. Her background in machine learning, multimodal systems, and reproducible experimentation aligns closely with advanced study in artificial intelligence. I recommend her with confidence and expect that she will contribute thoughtful research, dependable work, and a genuine commitment to improvement. Please feel free to contact me if further information would be helpful.",
                    "Sincerely,\n\n[Recommender Name]\n[Academic Title]\n[Department and Institution]\n[Telephone]\n[Email Address]\n[Signature]",
                ],
            )
        ],
    },
    {
        "filename": "04_研究计划_RP_导出模板.docx",
        "title": "RESEARCH PROPOSAL",
        "short": "RP",
        "document_type": "Research Proposal (RP)",
        "body_heading": "RESEARCH PROPOSAL",
        "mode": "sectioned",
        "sections": [
            (
                "PROPOSED TITLE",
                [
                    "Robust Multimodal Representation Alignment for Medical Image Analysis under Limited Supervision"
                ],
            ),
            (
                "ABSTRACT",
                [
                    "Multimodal learning can combine medical images with associated clinical text, but existing systems often depend on abundant paired data and degrade when annotations are limited or one modality is incomplete. This project proposes a robust representation-alignment framework that combines contrastive learning, prototype-based regularisation, and cross-modal attention. The study will evaluate whether the proposed approach improves alignment quality and downstream classification under low-sample and missing-modality conditions. Experiments will use public medical datasets, reproducible baselines, and predefined evaluation metrics. The expected contribution is a practical and carefully evaluated method for learning reliable multimodal representations in data-constrained clinical settings."
                ],
            ),
            (
                "RESEARCH BACKGROUND AND PROBLEM STATEMENT",
                [
                    "Medical decisions frequently draw on several forms of evidence, including images, reports, and structured patient information. Multimodal models offer a way to integrate these sources, yet their performance is often reported under conditions that assume well-paired, representative, and sufficiently large training data. Clinical datasets rarely satisfy all of these assumptions. Annotation is expensive, institutional data distributions differ, and text may be incomplete or only weakly related to the corresponding image.",
                    "My undergraduate project on deep-learning-based medical image analysis revealed that standard image-text contrastive models became unstable as labelled samples were reduced. This motivates a focused research problem: how can multimodal representations preserve useful cross-modal alignment and downstream performance when supervision is limited and input modalities are incomplete? The project will concentrate on image-text learning rather than general multimodal intelligence so that the proposed hypotheses can be tested within a realistic graduate research schedule."
                ],
            ),
            (
                "RESEARCH QUESTIONS",
                [
                    "The study will address three questions. First, can prototype-based regularisation stabilise contrastive image-text alignment when only a limited proportion of paired samples is available? Second, can an attention-based fusion mechanism reduce performance loss when text or image features are incomplete? Third, which components of the framework contribute most consistently across datasets, sample sizes, and evaluation tasks?"
                ],
            ),
            (
                "METHODOLOGY",
                [
                    "The research will begin with a reproducible implementation of established image-text contrastive baselines. Publicly available medical image and report datasets will be divided using fixed training, validation, and test protocols. Low-supervision settings will be simulated by controlling the proportion of paired samples, while missing-modality settings will be created through documented masking procedures. Data governance requirements and dataset licences will be reviewed before any experiment is conducted.",
                    "The proposed model will combine three components: modality-specific encoders, a contrastive alignment objective, and prototype-based regularisation that encourages samples with related semantic content to occupy stable regions of the shared representation space. A cross-modal attention module will then integrate available features while accounting for absent or uncertain inputs. Performance will be compared with unimodal and multimodal baselines using classification quality, retrieval quality, calibration, and robustness measures. Ablation studies will isolate the contribution of each component, and repeated runs will be reported with confidence intervals rather than a single best score."
                ],
            ),
            (
                "EXPECTED OUTCOMES AND CONTRIBUTION",
                [
                    "The anticipated outcome is a validated framework that improves multimodal alignment and limits performance degradation under constrained supervision. The academic contribution will be an analysis of when prototype regularisation and attention-based fusion are useful, including documented failure conditions. Practical outputs will include reproducible training code, evaluation protocols, ablation results, and a manuscript suitable for submission to a peer-reviewed machine-learning or medical-imaging venue. Any performance target will be treated as an evaluation objective rather than a guaranteed result."
                ],
            ),
            (
                "WORK PLAN",
                [
                    "Months 1–3 will cover the literature review, dataset governance checks, environment preparation, and baseline reproduction. Months 4–6 will focus on model design and initial controlled experiments. Months 7–9 will cover robustness evaluation, ablation studies, and error analysis. Months 10–12 will be used for consolidation, reproducibility checks, thesis or paper writing, and revision. Time will be reserved at the end of each phase to revise assumptions if the evidence does not support the original design."
                ],
            ),
            (
                "REFERENCES",
                [
                    "A complete reference list should be inserted automatically from the verified bibliography used in the final document. References should prioritise peer-reviewed work from the preceding three years while retaining foundational sources required to explain the selected methods."
                ],
            ),
        ],
    },
    {
        "filename": "05_套磁信_CL_导出模板.docx",
        "title": "COLD EMAIL",
        "short": "CL",
        "document_type": "Cold Email (CL)",
        "body_heading": "EMAIL CONTENT",
        "mode": "letter",
        "sections": [
            (
                None,
                [
                    "Subject: Prospective Master's Student — Multimodal Learning under Limited Supervision",
                    "Dear Professor [Surname],",
                    "I am Li Xue, a final-year Computer Science student at Peking University applying to Carnegie Mellon University's Artificial Intelligence programme for Fall 2026. Your recent work on robust multimodal alignment, particularly its treatment of distribution shifts, closely connects with a problem I encountered while leading a medical image-analysis project.",
                    "In that project, I investigated image-text alignment when labelled pairs were limited and tested contrastive and attention-based approaches in PyTorch. I also completed an NLP internship at Baidu AI Lab, where my work helped improve intent-recognition accuracy from 91.2 percent to 94.6 percent. These experiences have led me to study reliable multimodal representations in data-constrained settings.",
                    "I would be grateful for the opportunity to learn whether my background could fit your group's upcoming projects. I have attached my CV and transcript for context. If convenient, would you be available for a brief 15–30 minute online conversation during the week of [date range]? I am happy to adapt to your schedule.",
                    "Thank you for your time and consideration.",
                    "Best regards,\nLi Xue\nPeking University\n[Email Address] | [Telephone]",
                ],
            )
        ],
    },
]


def set_run_font(run, name="Times New Roman", size=11, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)
    set_run_font(run, name="Arial", size=8.5, color=(110, 110, 110))


def configure_document(document, short_name):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.footer_distance = Cm(1.1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    prefix = footer.add_run(f"Li Xue  |  {short_name}  |  Page ")
    set_run_font(prefix, name="Arial", size=8.5, color=(110, 110, 110))
    add_page_number(footer)


def add_title(document, title, short_name):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(title)
    set_run_font(run, name="Arial", size=18, bold=True, color=(31, 41, 55))

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.keep_with_next = True
    run = subtitle.add_run(f"Application Document — {short_name}")
    set_run_font(run, name="Arial", size=9.5, color=(100, 100, 100))


def add_section_heading(document, text, top_space=12, bottom_space=6):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(top_space)
    paragraph.paragraph_format.space_after = Pt(bottom_space)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, name="Arial", size=11.5, bold=True, color=(31, 41, 55))
    return paragraph


def add_metadata(document, document_type):
    add_section_heading(document, "STUDENT & APPLICATION INFORMATION", top_space=0, bottom_space=7)
    fields = [("Document Type", document_type)] + STUDENT_APPLICATION_INFO
    for label, value in fields:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_together = True
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, name="Arial", size=9.5, bold=True, color=(55, 65, 81))
        value_run = paragraph.add_run(value)
        set_run_font(value_run, name="Arial", size=9.5, color=(55, 65, 81))

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.line_spacing = 0.5


def add_body_paragraph(document, text, mode):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.widow_control = True

    if mode == "essay":
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    else:
        paragraph.paragraph_format.first_line_indent = Cm(0)

    lines = text.split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Times New Roman", size=11, color=(25, 25, 25))
    return paragraph


def add_document_body(document, definition):
    add_section_heading(document, definition["body_heading"], top_space=3, bottom_space=10)
    mode = definition["mode"]
    for section_title, paragraphs in definition["sections"]:
        if section_title:
            add_section_heading(document, section_title, top_space=10, bottom_space=5)
        for text in paragraphs:
            add_body_paragraph(document, text, mode)


def build_document(definition):
    document = Document()
    configure_document(document, definition["short"])
    add_title(document, definition["title"], definition["short"])
    add_metadata(document, definition["document_type"])
    add_document_body(document, definition)

    document.core_properties.title = definition["title"].title()
    document.core_properties.subject = "PBS study-abroad application document export template"
    document.core_properties.author = "PBS Study Abroad Application Management System"
    document.core_properties.keywords = definition["short"]

    output_path = OUTPUT_DIR / definition["filename"]
    document.save(output_path)
    return output_path


def main():
    generated = [build_document(definition) for definition in DOCUMENTS]
    print("Generated Word templates:")
    for path in generated:
        print(f"- {path.name}")


if __name__ == "__main__":
    main()
